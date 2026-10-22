"""Offer files and settlements for QEnergy (José Antonio), the Market Participant that submits our plants' offers to CENACE.

Flow: the Pantheon engine POSTs each day's hourly forecast → QEnergy downloads the sales offer ("Oferta No Despachable",
one day) and the purchase offer ("Oferta de Compra de Energía", up to 7 days) in QEnergy's own templates → after the day,
cleared MW, metered MWh and prices are posted → the settlement report shows day-ahead revenue and real-time deviations.
Pantheon never submits to CENACE itself.
"""
from __future__ import annotations

import csv
import io
import os
from datetime import date, datetime, timedelta, timezone
from functools import lru_cache
from typing import Literal, Optional

from fastapi import APIRouter, BackgroundTasks, Depends, Query
from fastapi.responses import Response
from pydantic import BaseModel, Field, field_validator

from ..auth import Caller, need
from ..errors import ApiError
from ..qenergy_store import SITES, QStore

router = APIRouter(prefix="/v1/qenergy", tags=["QEnergy offers and settlements"])
TEMPLATES = os.path.join(os.path.dirname(__file__), "..", "templates", "qenergy")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@lru_cache
def get_qstore() -> QStore:
    return QStore()


# ---------- schemas ----------
class Segment(BaseModel):
    mw: float = Field(ge=0)
    price: float = Field(description="MXN/MWh")


class ForecastHour(BaseModel):
    hour: int = Field(ge=1, le=25, description="CENACE hour 1–24 (25 only on the day clocks go back)")
    sell_mw: float = Field(ge=0, description="Energy offered for sale, MW (non-dispatchable offer)")
    sell_price: float = Field(0, description="MXN/MWh. 0 = price-taker, as in QEnergy's samples")
    buy_fixed_mw: float = Field(0, ge=0, description="Fixed purchase demand at the load centre, MW")
    buy_segments: list[Segment] = Field(default_factory=list, max_length=3, description="Up to 3 price-sensitive purchase blocks")


class ForecastInput(BaseModel):
    source: str = Field(max_length=120, description="e.g. 'CENACE PML MDA', 'GFS 0.25 run 00Z', 'plant SCADA'")
    as_of: datetime = Field(description="Latest moment the data covers / was published, with time zone")
    retrieved_at: datetime = Field(description="When the engine downloaded it, with time zone")
    sha256: str = Field(pattern=r"^[0-9a-f]{64}$", description="SHA-256 of the input file exactly as used")


class ForecastIn(BaseModel):
    hours: list[ForecastHour] = Field(min_length=24, max_length=25)
    model_version: str = Field("unknown", max_length=80)
    model_commit: Optional[str] = Field(None, pattern=r"^[0-9a-f]{40}$",
                                        description="Git commit (SHA-1) of the code that ran. Required once the plant has a Víspera chain")
    model_repo: Optional[str] = Field(None, max_length=200, description="Company repository holding that commit")
    inputs: list[ForecastInput] = Field(default_factory=list, max_length=50,
                                        description="Every input the model used. Required (at least one) once the plant has a Víspera chain")

    @field_validator("hours")
    @classmethod
    def one_per_hour(cls, v: list[ForecastHour]) -> list[ForecastHour]:
        hs = [h.hour for h in v]
        if sorted(hs) != list(range(1, len(v) + 1)):
            raise ValueError("hours must be 1..24 (or 1..25), each once")
        return sorted(v, key=lambda h: h.hour)


class ForecastReceipt(BaseModel):
    site: str
    day: date
    version: int
    sha256: str
    received_at: str
    model_version: str
    sample: bool = False
    chain: Optional[dict] = Field(None, description="The Víspera chain entry this forecast became: {chain, seq, hash}. "
                                  "None if the plant has no chain yet or the day is before its start date")


class ForecastOut(ForecastReceipt):
    posted_by: str
    hours: list[ForecastHour]


class DayListing(BaseModel):
    day: date
    versions: int
    last_received_at: str
    sample: bool


class ResultHour(BaseModel):
    hour: int = Field(ge=1, le=25)
    cleared_mw: float = Field(description="Assigned by CENACE in the day-ahead market")
    measured_mwh: float = Field(description="From the certified meter")
    pml_mda: float = Field(description="Day-ahead nodal price, MXN/MWh")
    pml_mtr: float = Field(description="Real-time nodal price, MXN/MWh")


class ResultsIn(BaseModel):
    hours: list[ResultHour] = Field(min_length=24, max_length=25)
    source: str = Field("CENACE", description="Where the prices came from")
    vintage: Optional[str] = Field(None, description="Which CENACE publication, e.g. 'MDA 2026-10-21 v1'")
    retrieved_at: Optional[str] = Field(None, description="When the prices were downloaded, ISO 8601")


class Site(BaseModel):
    id: str
    name: str
    central: str
    unidad: str
    centro_carga: str
    market_participant: str


def site_or_404(site_id: str) -> dict:
    s = SITES.get(site_id)
    if not s:
        raise ApiError(404, "not_found", what=site_id)
    return s


def forecast_or_404(store: QStore, site_id: str, day: date, version: Optional[int] = None) -> dict:
    f = store.forecast(site_id, day, version)
    if not f:
        raise ApiError(404, "not_found", what=f"forecast {site_id} {day}")
    return f


# ---------- endpoints ----------
@router.get("/sites", response_model=list[Site], summary="Plants QEnergy represents")
def sites(_: Caller = Depends(need("qenergy.read"))) -> list[Site]:
    return [Site(**s) for s in SITES.values()]


@router.post("/sites/{site_id}/forecasts/{day}", response_model=ForecastReceipt, status_code=201,
             summary="Pantheon engine posts a day's hourly forecast (new version each time)")
def post_forecast(site_id: str, day: date, body: ForecastIn, tasks: BackgroundTasks, c: Caller = Depends(need("forecast.write")),
                  store: QStore = Depends(get_qstore)) -> ForecastReceipt:
    site_or_404(site_id)
    ch = store.chain(site_id)
    chained = bool(ch) and day.isoformat() >= ch["terms"]["start_date"]
    if chained:
        problems = []
        if not body.model_commit or not body.model_repo:
            problems.append({"field": "model_commit", "message": "the git commit and repository of the code that ran are required on a Víspera chain"})
        if not body.inputs:
            problems.append({"field": "inputs", "message": "list every input (source, as_of, retrieved_at, sha256) on a Víspera chain"})
        now = datetime.now(timezone.utc)
        for i, x in enumerate(body.inputs):
            if x.as_of.tzinfo is None or x.retrieved_at.tzinfo is None:
                problems.append({"field": f"inputs[{i}]", "message": "as_of and retrieved_at need a time zone"})
            elif x.as_of > x.retrieved_at or x.retrieved_at > now:
                problems.append({"field": f"inputs[{i}]", "message": "as_of must be before retrieved_at, and both in the past"})
        if problems:
            raise ApiError(422, "validation_error", details=problems)
    payload = {"hours": [h.model_dump() for h in body.hours]}
    rec = store.put_forecast(site_id, day, payload, posted_by=c.actor, model_version=body.model_version)
    sealed = {"site": site_id, "model_version": body.model_version, "model_commit": body.model_commit, "model_repo": body.model_repo,
              "inputs": [x.model_dump(mode="json") for x in body.inputs], "record_sha256": rec["sha256"], **payload}
    rec["chain"] = to_chain(store, tasks, site_id, day, "forecast", sealed, rec["id"])
    return ForecastReceipt(**rec)


def to_chain(store: QStore, tasks: BackgroundTasks, chain_id: str, day: date, kind: str, payload: dict, forecast_id: Optional[int] = None) -> Optional[dict]:
    """Append to the plant's Víspera chain if it has one and the day is on or after its start date; anchor after the reply."""
    ch = store.chain(chain_id)
    if not ch or day.isoformat() < ch["terms"]["start_date"]:
        return None
    from .vispera import anchor_entry
    e = store.append(chain_id, kind, day.isoformat(), payload, forecast_id)
    tasks.add_task(anchor_entry, store, e["id"])
    return {"chain": chain_id, "seq": e["seq"], "hash": e["hash"]}


def _chain_ref(f: dict) -> Optional[dict]:
    e = f.get("entry")
    return {"chain": e["chain"], "seq": e["seq"], "hash": e["hash"]} if e else None


@router.get("/sites/{site_id}/forecasts", response_model=list[DayListing], summary="Days with a forecast, newest first")
def list_forecasts(site_id: str, _: Caller = Depends(need("qenergy.read")), store: QStore = Depends(get_qstore)) -> list[DayListing]:
    site_or_404(site_id)
    return [DayListing(**d) for d in store.forecast_days(site_id)]


@router.get("/sites/{site_id}/forecasts/{day}", response_model=ForecastOut, summary="Latest forecast for a day, with its fingerprint")
def get_forecast(site_id: str, day: date, _: Caller = Depends(need("qenergy.read")), store: QStore = Depends(get_qstore)) -> ForecastOut:
    site_or_404(site_id)
    f = forecast_or_404(store, site_id, day)
    return ForecastOut(site=site_id, day=day, version=f["version"], sha256=f["sha256"], received_at=f["received_at"],
                       model_version=f["model_version"], sample=f["sample"], posted_by=f["posted_by"], hours=f["payload"]["hours"],
                       chain=_chain_ref(f))


@router.get("/sites/{site_id}/forecasts/{day}/canonical", summary="The exact bytes that were fingerprinted, to re-check the SHA-256 independently")
def forecast_canonical(site_id: str, day: date, version: Optional[int] = Query(None, ge=1, description="Default: latest"),
                       _: Caller = Depends(need("qenergy.read")), store: QStore = Depends(get_qstore)) -> Response:
    site_or_404(site_id)
    f = forecast_or_404(store, site_id, day, version)
    return Response(f["canonical"], media_type="application/json", headers={"X-Forecast-SHA256": f["sha256"]})


@router.get("/sites/{site_id}/forecasts/{day}/history", summary="Every version posted for a day (time and SHA-256)")
def forecast_history(site_id: str, day: date, _: Caller = Depends(need("qenergy.read")), store: QStore = Depends(get_qstore)) -> list[dict]:
    site_or_404(site_id)
    return store.history(site_id, day)


def _xlsx(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _load(name: str):
    try:
        import openpyxl
    except ImportError:  # pragma: no cover
        raise ApiError(501, "xlsx_unavailable")
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(os.path.join(TEMPLATES, name))


@router.get("/sites/{site_id}/offers/{day}/venta.xlsx", summary="Sales offer (Oferta No Despachable) for one day, in QEnergy's template")
def offer_venta(site_id: str, day: date, _: Caller = Depends(need("qenergy.read")), store: QStore = Depends(get_qstore)) -> Response:
    s = site_or_404(site_id)
    f = forecast_or_404(store, site_id, day)
    wb = _load("oferta_venta.xlsx")
    ws = wb.active
    when = datetime(day.year, day.month, day.day)
    ws["D4"], ws["H4"], ws["L4"], ws["P4"] = when, when, s["central"], s["unidad"]
    by_hour = {h["hour"]: h for h in f["payload"]["hours"]}
    for hour in range(1, 26):
        r = 13 + hour
        h = by_hour.get(hour)
        ws[f"D{r}"] = round(h["sell_mw"], 6) if h else (None if hour == 25 else 0)
        ws[f"F{r}"] = h["sell_price"] if h else (None if hour == 25 else 0)
    name = f"{s['central']}_OfertaVenta_{day.isoformat()}.xlsx"
    return Response(_xlsx(wb), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"', "X-Forecast-SHA256": f["sha256"]})


@router.get("/sites/{site_id}/offers/{day}/compra.xlsx", summary="Purchase offer (Oferta de Compra) from a day, same hourly profile for 1–7 days")
def offer_compra(site_id: str, day: date, days: int = Query(7, ge=1, le=7), _: Caller = Depends(need("qenergy.read")),
                 store: QStore = Depends(get_qstore)) -> Response:
    s = site_or_404(site_id)
    f = forecast_or_404(store, site_id, day)
    wb = _load("oferta_compra.xlsx")
    ws = wb.active
    ws["D3"] = datetime(day.year, day.month, day.day)
    ws["H3"] = "=D3+6" if days == 7 else datetime(day.year, day.month, day.day) + timedelta(days=days - 1)
    ws["M3"] = s["centro_carga"]
    by_hour = {h["hour"]: h for h in f["payload"]["hours"]}
    cols = [("F", "H"), ("J", "L"), ("N", "P")]
    for hour in range(1, 26):
        r = 13 + hour
        h = by_hour.get(hour)
        if not h and hour == 25:
            for c in ["D"] + [x for p in cols for x in p]:
                ws[f"{c}{r}"] = None
            continue
        ws[f"D{r}"] = round(h["buy_fixed_mw"], 6) if h else 0
        segs = (h["buy_segments"] if h else []) + [{"mw": 0, "price": 0}] * 3
        for (mw_c, pr_c), seg in zip(cols, segs[:3]):
            ws[f"{mw_c}{r}"] = round(seg["mw"], 6)
            ws[f"{pr_c}{r}"] = seg["price"]
    name = f"{s['centro_carga']}_OfertaCompra_{day.isoformat()}_{days}d.xlsx"
    return Response(_xlsx(wb), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"', "X-Forecast-SHA256": f["sha256"]})


@router.post("/sites/{site_id}/results/{day}", status_code=201, summary="Post a day's cleared MW, metered MWh and nodal prices")
def post_results(site_id: str, day: date, body: ResultsIn, tasks: BackgroundTasks, c: Caller = Depends(need("qenergy.write")),
                 store: QStore = Depends(get_qstore)) -> dict:
    site_or_404(site_id)
    store.put_results(site_id, day, {"hours": [h.model_dump() for h in body.hours]}, posted_by=c.actor)
    # Only the public prices go into the chain (in clear), never QEnergy's cleared or metered volumes.
    actuals = {"site": site_id, "source": body.source, "vintage": body.vintage, "retrieved_at": body.retrieved_at,
               "prices": [{"hour": h.hour, "pml_mda": h.pml_mda, "pml_mtr": h.pml_mtr} for h in body.hours]}
    ref = to_chain(store, tasks, site_id, day, "actuals", actuals)
    return {"site": site_id, "day": day.isoformat(), "status": "stored", "chain": ref}


def _settlement_rows(store: QStore, site_id: str, month: str) -> tuple[list[dict], list[dict], bool]:
    res = store.results_month(site_id, month)
    hours, days, sample = [], [], False
    for d, payload in res.items():
        f = store.forecast(site_id, date.fromisoformat(d))
        sample = sample or (f is not None and f["sample"])
        offered = {h["hour"]: h["sell_mw"] for h in (f["payload"]["hours"] if f else [])}
        tot = {"day": d, "offered_mwh": 0.0, "cleared_mwh": 0.0, "measured_mwh": 0.0, "mda_mxn": 0.0, "mtr_mxn": 0.0}
        for h in payload["hours"]:
            mda = h["cleared_mw"] * h["pml_mda"]
            mtr = (h["measured_mwh"] - h["cleared_mw"]) * h["pml_mtr"]
            row = {"day": d, "hour": h["hour"], "offered_mw": offered.get(h["hour"], 0.0), "cleared_mw": h["cleared_mw"],
                   "measured_mwh": h["measured_mwh"], "pml_mda": h["pml_mda"], "pml_mtr": h["pml_mtr"],
                   "mda_mxn": round(mda, 2), "mtr_mxn": round(mtr, 2), "total_mxn": round(mda + mtr, 2)}
            hours.append(row)
            tot["offered_mwh"] += row["offered_mw"]; tot["cleared_mwh"] += h["cleared_mw"]; tot["measured_mwh"] += h["measured_mwh"]
            tot["mda_mxn"] += mda; tot["mtr_mxn"] += mtr
        tot["total_mxn"] = tot["mda_mxn"] + tot["mtr_mxn"]
        days.append({k: (round(v, 4) if isinstance(v, float) else v) for k, v in tot.items()})
    return hours, days, sample


@router.get("/sites/{site_id}/settlements", summary="Settlement for a month: day-ahead revenue + real-time deviations, by day and hour")
def settlements(site_id: str, month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), format: Literal["json", "csv", "xlsx"] = "json",
                _: Caller = Depends(need("qenergy.read")), store: QStore = Depends(get_qstore)):
    s = site_or_404(site_id)
    hours, days, sample = _settlement_rows(store, site_id, month)
    if format == "json":
        return {"site": site_id, "month": month, "sample_data": sample, "days": days, "hours": hours,
                "method": "MDA = cleared MW × PML MDA; MTR = (metered MWh − cleared MW) × PML MTR; total = MDA + MTR. MXN."}
    cols = ["day", "hour", "offered_mw", "cleared_mw", "measured_mwh", "pml_mda", "pml_mtr", "mda_mxn", "mtr_mxn", "total_mxn"]
    if format == "csv":
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=cols)
        w.writeheader()
        w.writerows(hours)
        return Response(buf.getvalue(), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{s["central"]}_Liquidacion_{month}.csv"'})
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append([f"Liquidación {s['name']} ({s['central']}) — {month}" + (" — DATOS DE EJEMPLO" if sample else "")])
    ws.append(["MDA = MW asignados × PML MDA; MTR = (MWh medidos − MW asignados) × PML MTR. MXN."])
    head = ["Día", "Ofertado MWh", "Asignado MWh", "Medido MWh", "Ingreso MDA", "Desviación MTR", "Total"]
    ws.append(head)
    for d in days:
        ws.append([d["day"], d["offered_mwh"], d["cleared_mwh"], d["measured_mwh"], round(d["mda_mxn"], 2), round(d["mtr_mxn"], 2), round(d["total_mxn"], 2)])
    if days:
        n = len(days)
        ws.append(["Total"] + [f"=SUM({c}4:{c}{3 + n})" for c in "BCDEFG"])
    for c in ws[1] + ws[3]:
        c.font = Font(bold=True)
    hs = wb.create_sheet("Horas")
    hs.append(["Día", "Hora", "Ofertado MW", "Asignado MW", "Medido MWh", "PML MDA", "PML MTR", "Ingreso MDA", "Desviación MTR", "Total"])
    for c in hs[1]:
        c.font = Font(bold=True)
    for h in hours:
        hs.append([h[k] for k in cols])
    return Response(_xlsx(wb), media_type=XLSX,
                    headers={"Content-Disposition": f'attachment; filename="{s["central"]}_Liquidacion_{month}.xlsx"'})
