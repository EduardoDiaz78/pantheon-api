import hashlib
import hmac
import json
from datetime import datetime

import pytest
from fastapi.testclient import TestClient

from app.auth import get_now, get_store
from app.main import app
from app.planning import MX_TZ
from app.store import Store

DAY = "2026-10-18"
MORNING = datetime(2026, 10, 17, 8, 15, tzinfo=MX_TZ)
LATE = datetime(2026, 10, 17, 10, 1, tzinfo=MX_TZ)
EVENING = datetime(2026, 10, 17, 19, 42, tzinfo=MX_TZ)

VIEWER, AUDITOR, OPERATOR, MANAGER, ADMIN = ({"X-API-Key": k} for k in
                                            ("demo-viewer", "demo-auditor", "demo-operator", "demo-manager", "demo-admin"))
ES = {"Accept-Language": "es-MX"}


@pytest.fixture
def store():
    return Store()


@pytest.fixture
def client(store):
    app.dependency_overrides[get_store] = lambda: store
    app.dependency_overrides[get_now] = lambda: MORNING
    yield TestClient(app)
    app.dependency_overrides.clear()


def at(t):
    app.dependency_overrides[get_now] = lambda: t


def plan(c, site="site-a", headers=VIEWER, **q):
    r = c.get(f"/v1/sites/{site}/plans/{DAY}", headers=headers, params=q)
    assert r.status_code == 200, r.text
    return r.json()


# ---------- the numbers match the screens ----------

def test_plan_matches_today_screen(client):
    p = plan(client)
    s = p["summary"]
    assert (s["energy_margin"], s["discharge_revenue"], s["charging_cost"], s["wear_cost"]) == (73495, 99080, 19465, 6120)
    assert (s["solar_total_mwh"], s["solar_to_battery_mwh"], s["solar_to_grid_mwh"]) == (225, 34, 191)
    assert s["cycles"] == 0.85 and s["likely_critical_hours"] == 3 and s["likely_critical_hours_covered"] == 3
    assert p["status"] == "draft" and p["minutes_to_deadline"] == 105 and p["limit_check"]["ok"]
    assert [h["hour"] for h in p["hours"] if h["likely_critical"]] == [19, 20, 21]
    assert p["hours"][20]["critical_likelihood"] == 0.81
    assert p["hours"][12]["charge_from_solar_mw"] == 8.5 and p["hours"][12]["charge_from_grid_mw"] == 0
    assert [(b["start"], b["end"], b["action"]) for b in p["schedule"]] == [("11:00", "15:00", "charge"), ("18:00", "23:00", "discharge")]
    assert p["schedule"][0]["source"] == "Charge from plant solar"


def test_strategy_cards_match_screen(client):
    opts = {o["strategy"]: o for o in client.get(f"/v1/sites/site-a/plans/{DAY}/strategies", headers=VIEWER).json()}
    assert opts["conservative"]["energy_margin"] == 66005 and opts["conservative"]["cycles"] == 0.75
    assert opts["balanced"]["energy_margin"] == 73495
    assert opts["aggressive"]["energy_margin"] == 72775 and opts["aggressive"]["cycles"] == 1.05
    assert opts["aggressive"]["energy_margin"] < opts["balanced"]["energy_margin"]
    assert not opts["aggressive"]["within_limits"]
    assert [v["code"] for v in opts["aggressive"]["violations"]] == ["cycles_per_day"]


def test_capacity_is_always_a_projection(client):
    cap = plan(client)["capacity"]
    assert cap["status"] == "projection" and cap["annual_value_mxn"] is None
    assert cap["battery_available_mw_in_flagged_hours"] == 8
    wf = client.get("/v1/sites/site-a/settlements?month=2026-10", headers=VIEWER).json()["waterfall"]
    assert wf[-1] == {"label": "Capacity payments", "amount": None, "kind": "pending"}
    sim = client.post("/v1/simulations", headers=VIEWER, json={"region": "Nuevo León", "power_mw": 20, "duration_h": 4}).json()
    assert sim["capacity_value"] == "not_included"


def test_every_operating_site_plan_is_within_limits(client):
    for sid in ("site-a", "site-b", "site-c", "site-d", "site-e"):
        assert plan(client, sid)["limit_check"]["ok"], sid


# ---------- the prototype flow ----------

def test_demo_flow_end_to_end(client, store):
    home = client.get("/v1/home", headers=OPERATOR).json()
    first = home["actions"][0]
    assert first["kind"] == "approve_plan" and first["site_id"] == "site-a" and first["minutes_left"] == 105
    assert {a["site_id"] for a in home["actions"] if a["kind"] == "approve_plan"} == {"site-a", "site-c"}  # B, D, E already approved

    url = f"/v1/sites/site-a/plans/{DAY}/approve"
    assert client.post(url, json={"strategy": "aggressive"}, headers=OPERATOR).json()["error"]["code"] == "soft_needs_ack"
    assert client.post(url, json={"strategy": "aggressive", "acknowledge_soft_violations": True}, headers=OPERATOR).status_code == 403
    r = client.post(url, json={"strategy": "balanced", "note": "Inverter check moved to Tuesday."}, headers=OPERATOR)
    assert r.status_code == 200 and r.json()["status"] == "approved" and r.json()["approved_by"] == "Daniela"

    al1 = client.get("/v1/alerts/al-1", headers=VIEWER).json()
    assert al1["status"] == "resolved" and al1["resolution"] == "plan_approved"
    assert client.get(f"/v1/sites/site-a/plans/{DAY}/offer-file", headers=OPERATOR).text.count("\n") == 26
    log = client.get("/v1/audit-log", headers=MANAGER).json()
    assert log[0]["action"] == "plan.export" and log[1]["action"] == "plan.approve"
    assert "Inverter check" in log[1]["detail"]


def test_manager_signs_off_aggressive(client):
    r = client.post(f"/v1/sites/site-a/plans/{DAY}/approve", headers=MANAGER,
                    json={"strategy": "aggressive", "acknowledge_soft_violations": True})
    assert r.status_code == 200 and r.json()["strategy"] == "aggressive"
    assert "cycles_per_day" in client.get("/v1/audit-log", headers=MANAGER).json()[0]["detail"]


def test_hard_limits_never_approvable(client):
    r = client.post(f"/v1/sites/site-a/plans/{DAY}/overrides", headers=OPERATOR,
                    json={"hour": 20, "dispatch_mw": 15, "reason": "push it"})
    assert {v["code"] for v in r.json()["limit_check"]["violations"]} >= {"discharge_power"}
    bad = client.post(f"/v1/sites/site-a/plans/{DAY}/approve", headers=ADMIN, json={"acknowledge_soft_violations": True})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "hard_violations"
    client.delete(f"/v1/sites/site-a/plans/{DAY}/overrides/20", headers=OPERATOR)
    assert client.post(f"/v1/sites/site-a/plans/{DAY}/approve", headers=OPERATOR, json={}).status_code == 200


# ---------- states ----------

def test_deadline_passed_state(client):
    at(LATE)
    p = plan(client)
    assert p["status"] == "missed" and p["minutes_to_deadline"] == 0
    r = client.post(f"/v1/sites/site-a/plans/{DAY}/approve", headers=OPERATOR, json={})
    assert r.status_code == 409 and r.json()["error"]["code"] == "deadline_passed"
    assert not [a for a in client.get("/v1/home", headers=OPERATOR).json()["actions"] if a["kind"] == "approve_plan"]


def test_onboarding_and_maintenance_states(client):
    r = client.get(f"/v1/sites/site-g/plans/{DAY}", headers=VIEWER)
    assert r.status_code == 409 and r.json()["error"]["code"] == "no_data_yet" and "2026-10-20" in r.json()["error"]["message"]
    assert client.get(f"/v1/sites/site-f/plans/{DAY}", headers=VIEWER).json()["error"]["code"] == "site_maintenance"


def test_stale_data_state_and_incident_decision(client, store):
    feeds = {f["name"]: f for f in plan(client, "site-c")["data_status"]}
    assert feeds["site_telemetry"]["status"] == "stale" and feeds["site_telemetry"]["age_minutes"] == 33
    assert any(a["kind"] == "incident" for a in client.get("/v1/home", headers=OPERATOR).json()["actions"])
    bad = client.post("/v1/alerts/al-3/resolve", headers=OPERATOR, json={"choice": "acknowledge"})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "bad_decision"
    assert client.post("/v1/alerts/al-3/resolve", headers=VIEWER, json={"choice": "hold_battery"}).status_code == 403
    r = client.post("/v1/alerts/al-3/resolve", headers=OPERATOR, json={"choice": "hold_battery", "note": "Wait for data"})
    assert r.json()["status"] == "resolved" and store.sites["site-c"].status == "holding"
    assert client.post("/v1/alerts/al-3/resolve", headers=OPERATOR, json={"choice": "keep_plan"}).json()["error"]["code"] == "alert_closed"


def test_escalation(client):
    r = client.post("/v1/alerts/al-2/escalate", headers=OPERATOR, json={"to_user_id": "u-mauricio", "note": "Please look"})
    assert r.json()["escalated_to"] == "Mauricio"


def test_live_plant_view_during_critical_hour(client):
    at(EVENING)
    n = client.get("/v1/sites/site-a/now", headers=VIEWER).json()
    assert n["in_likely_critical_hour"] and n["battery_mw"] == 10 and n["solar_mw"] == 1
    assert n["soc_pct"] == 70.0 and n["soc_mwh"] == 28.0  # 35 MWh at 19:00, minus 42 min at 10 MW
    assert n["next_action"] == "Hold until tomorrow's plan" and n["on_plan"]
    assert client.get("/v1/sites/site-g/now", headers=VIEWER).status_code == 409


# ---------- roles ----------

def test_permission_matrix(client):
    assert client.get("/v1/audit-log", headers=AUDITOR).status_code == 200
    assert client.get("/v1/audit-log", headers=OPERATOR).status_code == 403
    assert client.get("/v1/sites/site-a/forecast-accuracy/hourly.csv", headers=AUDITOR).status_code == 200
    assert client.get("/v1/sites/site-a/forecast-accuracy/hourly.csv", headers=VIEWER).status_code == 403
    assert client.post(f"/v1/sites/site-a/plans/{DAY}/approve", headers=AUDITOR, json={}).status_code == 403
    assert client.put("/v1/sites/site-a/limits", headers=MANAGER, json={}).status_code == 403
    assert client.get("/v1", headers=VIEWER).status_code == 404
    me = client.get("/v1/me", headers=AUDITOR).json()
    assert me["user"]["role"] == "auditor" and "audit.read" in me["permissions"] and "plan.approve" not in me["permissions"]
    assert client.get("/v1/me").status_code == 401 and client.get("/v1/me").json()["error"]["code"] == "unauthorized"


def test_role_change_is_audited(client):
    assert client.put("/v1/users/u-lucia/role", headers=MANAGER, json={"role": "operator"}).status_code == 403
    r = client.put("/v1/users/u-lucia/role", headers=ADMIN, json={"role": "asset_manager"})
    assert r.json()["role"] == "asset_manager"
    assert client.get("/v1/audit-log", headers=ADMIN).json()[0]["detail"] == "Lucía: viewer → asset_manager"


# ---------- limits ----------

def test_limits_and_wear(client):
    url = "/v1/sites/site-a/limits"
    assert client.put(url, headers=ADMIN, json={"min_soc_pct": 50, "max_soc_pct": 40}).json()["error"]["code"] == "invalid_request"
    assert client.put(url, headers=ADMIN, json={"max_discharge_mw": 20}).json()["error"]["code"] == "over_rated"
    client.put(url, headers=ADMIN, json={"wear_cost_per_mwh": 0})
    assert plan(client)["summary"]["energy_margin"] == 73495 + 6120
    client.put(url, headers=ADMIN, json={"critical_flag_threshold": 0.3})
    p = plan(client)
    assert [h["hour"] for h in p["hours"] if h["likely_critical"]] == [18, 19, 20, 21]


# ---------- finance ----------

def test_settlement_dispute_and_report(client):
    s = client.get("/v1/sites/site-a/settlements?month=2026-10", headers=VIEWER).json()
    w = {x["label"]: x["amount"] for x in s["waterfall"]}
    assert w["Discharge revenue"] + w["Charging cost"] + w["Market & CENACE fees"] == pytest.approx(s["settled_total"])
    assert s["review_count"] == 1
    assert client.post("/v1/sites/site-a/disputes", headers=OPERATOR, json={"date": "2026-10-04", "reason": "Metered output higher"}).status_code == 403
    assert client.post("/v1/sites/site-a/disputes", headers=MANAGER, json={"date": "2026-10-09", "reason": "Not settled yet"}).json()["error"]["code"] == "day_not_settled"
    d = client.post("/v1/sites/site-a/disputes", headers=MANAGER, json={"date": "2026-10-04", "reason": "Metered output higher than settled"})
    assert d.status_code == 201 and d.json()["difference_pct"] == -5.85
    day4 = next(x for x in client.get("/v1/sites/site-a/settlements?month=2026-10", headers=VIEWER).json()["days"] if x["date"] == "2026-10-04")
    assert day4["status"] == "disputed"
    rep = client.get("/v1/sites/site-a/reports/monthly?month=2026-10", headers=VIEWER).json()
    assert rep["capacity"]["annual_value_mxn"] is None and rep["track_record"]["critical_hours_caught"] is None
    assert "7 days settled" in rep["headline"]


# ---------- language ----------

def test_spanish(client):
    p = plan(client, headers={**VIEWER, **ES})
    assert p["explanation"].startswith("La energía solar") and p["schedule"][0]["source"] == "Carga con la solar de la planta"
    opts = client.get(f"/v1/sites/site-a/plans/{DAY}/strategies", headers={**VIEWER, **ES}).json()
    assert opts[2]["violations"][0]["message"] == "1.05 ciclos planeados; el límite diario es 1."
    assert client.get("/v1/alerts", headers={**VIEWER, **ES}).json()[0]["title"] == "Aprueba el plan de mañana antes de las 10:00"
    assert client.get("/v1/me", headers=ES).json()["error"]["message"].startswith("Falta el encabezado")


# ---------- integrations ----------

def test_webhooks_signed_and_queued(client, store):
    assert client.post("/v1/webhooks", headers=MANAGER, json={"url": "https://example.com/hook", "events": ["plan.approved"]}).status_code == 403
    assert client.post("/v1/webhooks", headers=ADMIN, json={"url": "http://insecure", "events": ["plan.approved"]}).status_code == 422
    r = client.post("/v1/webhooks", headers=ADMIN, json={"url": "https://example.com/hook", "events": ["plan.approved", "alert.resolved"]}).json()
    secret, wid = r["secret"], r["id"]
    assert client.get("/v1/webhooks", headers=ADMIN).json()[0]["secret_hint"] == secret[-4:]
    client.post(f"/v1/sites/site-a/plans/{DAY}/approve", headers=OPERATOR, json={})
    dl = client.get(f"/v1/webhooks/{wid}/deliveries", headers=ADMIN).json()
    assert [d["event"] for d in dl] == ["plan.approved"] and dl[0]["payload"]["data"]["energy_margin"] == 73495
    body = json.dumps(dl[0]["payload"], sort_keys=True, ensure_ascii=False).encode()
    assert hmac.new(secret.encode(), body, hashlib.sha256).hexdigest() == dl[0]["signature"]


def test_simulation(client):
    r = client.post("/v1/simulations", headers=VIEWER, json={"region": "Nuevo León", "power_mw": 20, "duration_h": 4}).json()
    assert r["annual_net_p50"] == 58_000_000
    solar = client.post("/v1/simulations", headers=VIEWER, json={"region": "Nuevo León", "power_mw": 20, "duration_h": 4, "solar_mw": 60}).json()
    assert solar["annual_net_p50"] == round(58_000_000 * 1.06)
    assert client.post("/v1/simulations", headers=VIEWER, json={"region": "Nuevo León", "power_mw": 20, "duration_h": 3}).json()["error"]["code"] == "invalid_request"


def test_openapi(client):
    spec = client.get("/openapi.json").json()
    assert spec["info"]["version"] == "0.2.0" and len(spec["paths"]) >= 30



def test_demo_reset_restores_sample_data():
    app.dependency_overrides.clear()
    c = TestClient(app)
    assert c.post("/v1/alerts/al-3/resolve", json={"choice": "keep_plan"}, headers=OPERATOR).status_code == 200
    assert c.get("/v1/alerts/al-3", headers=OPERATOR).json()["status"] == "resolved"
    assert c.post("/v1/demo/reset", headers=OPERATOR).json() == {"status": "reset"}
    assert c.get("/v1/alerts/al-3", headers=OPERATOR).json()["status"] == "open"
    assert c.post("/v1/demo/reset").status_code == 401


def test_cors_allows_browser_calls():
    c = TestClient(app)
    r = c.options("/v1/home", headers={"Origin": "https://app.example.com", "Access-Control-Request-Method": "GET",
                                       "Access-Control-Request-Headers": "X-API-Key"})
    assert r.status_code == 200
    assert r.headers["access-control-allow-origin"] in ("*", "https://app.example.com")


# ---------- daily replay ----------

def test_replay_matches_replay_screen_and_home(client):
    r = client.get("/v1/sites/site-a/replay/2026-10-16", headers=VIEWER)
    assert r.status_code == 200, r.text
    j = r.json()
    s = j["summary"]
    assert (s["value_vs_idle"], s["plan_value"], s["vs_plan"], s["vs_plan_pct"]) == (71860, 70385, 1475, 2.1)
    assert s["critical_hours"] == [19, 20] and s["critical_hours_covered"] == "2 / 2" and s["fee_mxn"] is None
    assert [(d["code"], d["amount"]) for d in j["drivers"]] == [("prices", 2870), ("dispatch", -1395)]
    assert sum(d["amount"] for d in j["drivers"]) == s["vs_plan"]
    h = {x["hour"]: x for x in j["hours"]}
    assert (h[14]["planned_mw"], h[14]["actual_mw"]) == (-8.5, -7) and (h[22]["planned_mw"], h[22]["actual_mw"]) == (3, 1.5)
    assert h[20]["price_actual"] == 3380 and h[20]["value_vs_idle"] == 32000 and h[20]["confirmed_critical"]
    assert j["hours"][-1]["running_total"] == 71860 and j["earned_its_fee"] and j["headline"].startswith("Yes.")
    home = client.get("/v1/home", headers=VIEWER).json()
    assert (home["yesterday_energy_margin"], home["yesterday_vs_plan_pct"], home["yesterday_critical_hours"]) == (71860, 2.1, "2 / 2")


def test_replay_spanish_and_missing_day(client):
    j = client.get("/v1/sites/site-a/replay/2026-10-16", headers={**VIEWER, **ES}).json()
    assert j["headline"].startswith("Sí.")
    r = client.get("/v1/sites/site-a/replay/2026-10-15", headers=VIEWER)
    assert r.status_code == 404 and r.json()["error"]["code"] == "no_sample_day"


# ---------- WhatsApp ----------

DANIELA_WA = "+52 81 0000 4821"


def test_whatsapp_plan_message_and_numbers(client):
    m = client.get("/v1/channels/whatsapp/plan-message", params={"site_id": "site-a"}, headers=OPERATOR).json()
    assert m["to"] == "+52 ••• 4821" and m["context"] == {"site_id": "site-a", "date": DAY}
    assert [b["id"] for b in m["buttons"]] == ["approve", "see_details"] and m["can_approve_here"]
    assert "73,495 MXN" in m["text"] and "19:00–22:00" in m["text"]
    es = client.get("/v1/channels/whatsapp/plan-message", params={"site_id": "site-a"}, headers={**OPERATOR, **ES}).json()
    assert es["buttons"][0]["title"] == "Aprobar"
    assert client.get("/v1/me", headers=OPERATOR).json()["whatsapp_number"] == DANIELA_WA
    assert client.get("/v1/me", headers=VIEWER).json()["whatsapp_number"] is None
    assert client.get("/v1/channels/whatsapp/plan-message", params={"site_id": "site-a"}, headers=VIEWER).json()["to"] is None


def test_whatsapp_approve_is_audited_with_channel(client, store):
    r = client.post("/v1/channels/whatsapp/inbound", json={"from_number": DANIELA_WA, "button_id": "approve",
                                                           "context": {"site_id": "site-a", "date": DAY}})
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["status"] == "approved" and "Daniela · 08:15" in j["reply_text"]
    assert j["audit"]["action"] == "plan.approve" and "Channel: WhatsApp, verified number +52 ••• 4821" in j["audit"]["detail"]
    p = plan(client)
    assert p["status"] == "approved" and p["approved_by"] == "Daniela"
    again = client.post("/v1/channels/whatsapp/inbound", json={"from_number": DANIELA_WA, "button_id": "approve",
                                                               "context": {"site_id": "site-a", "date": DAY}})
    assert again.status_code == 409 and again.json()["error"]["code"] == "plan_locked"


def test_whatsapp_refuses_signoff_plans_and_unknown_numbers(client, store):
    from app.models import Strategy
    store.plan_state("site-a", datetime(2026, 10, 18).date()).strategy = Strategy.aggressive
    r = client.post("/v1/channels/whatsapp/inbound", json={"from_number": "+52 81 0000 7310", "button_id": "approve",
                                                           "context": {"site_id": "site-a", "date": DAY}})
    assert r.status_code == 409 and r.json()["error"]["code"] == "needs_app_signoff"
    m = client.get("/v1/channels/whatsapp/plan-message", params={"site_id": "site-a"}, headers=MANAGER).json()
    assert [b["id"] for b in m["buttons"]] == ["see_details"] and m["can_approve_here"] is False
    u = client.post("/v1/channels/whatsapp/inbound", json={"from_number": "+1 555 0100", "button_id": "approve",
                                                           "context": {"site_id": "site-a", "date": DAY}})
    assert u.status_code == 403 and u.json()["error"]["code"] == "wa_unknown_number"


def test_whatsapp_incident_buttons(client):
    m = client.get("/v1/channels/whatsapp/alert-message/al-3", headers=OPERATOR).json()
    assert [b["id"] for b in m["buttons"]] == ["keep_plan", "hold_battery", "hand_off"]
    r = client.post("/v1/channels/whatsapp/inbound", json={"from_number": DANIELA_WA, "button_id": "hand_off", "context": {"alert_id": "al-3"}})
    assert r.json()["status"] == "escalated" and "Mauricio" in r.json()["reply_text"]
    r = client.post("/v1/channels/whatsapp/inbound", json={"from_number": DANIELA_WA, "button_id": "hold_battery", "context": {"alert_id": "al-3"}})
    assert r.status_code == 200 and r.json()["status"] == "resolved" and "Channel: WhatsApp" in r.json()["audit"]["detail"]
    assert client.get("/v1/sites/site-c", headers=VIEWER).json()["status"] == "holding"
    d = client.post("/v1/channels/whatsapp/inbound", json={"from_number": DANIELA_WA, "button_id": "see_details",
                                                           "context": {"site_id": "site-a", "date": DAY}}).json()
    assert d["status"] == "link" and d["link"] == f"/v1/sites/site-a/plans/{DAY}"


# ---------------- CENACE: offers, cleared, imports, interconnection, battery state ----------------

def approve_a(client):
    r = client.post(f"/v1/sites/site-a/plans/{DAY}/approve", json={"strategy": "balanced"}, headers=OPERATOR)
    assert r.status_code == 200, r.text
    return r.json()


def test_offer_requires_approval_then_exports_trader_layout(client):
    assert client.get(f"/v1/sites/site-a/offer/{DAY}", headers=VIEWER).json()["error"]["code"] == "approve_first"
    approve_a(client)
    o = client.get(f"/v1/sites/site-a/offer/{DAY}", headers=VIEWER).json()
    assert o["status"] == "offered" and len(o["hours"]) == 24 and o["hours"][0]["offer_type"] == "self_schedule"
    assert o["deadline"].startswith("2026-10-17T10:00") and o["submitted_at"] is None
    h = o["hours"][12]
    assert abs(h["net_export_mw"] - (h["solar_mw"] + h["battery_mw"])) < 1e-6
    csv_ = client.get(f"/v1/sites/site-a/offer/{DAY}/export", headers=OPERATOR)
    assert csv_.status_code == 200 and csv_.text.splitlines()[0] == "HORA,MW_NETO,MW_BATERIA,MW_SOLAR,TIPO" and csv_.text.count("\n") == 25
    assert "AUTODESPACHO" in csv_.text and "attachment" in csv_.headers["content-disposition"]
    x = client.get(f"/v1/sites/site-a/offer/{DAY}/export?format=xlsx", headers=OPERATOR)
    assert x.status_code in (200, 501)
    assert client.get(f"/v1/sites/site-a/offer/{DAY}/export", headers=VIEWER).status_code == 403


def test_offer_submitted_and_deadline(client, store):
    approve_a(client)
    r = client.post(f"/v1/sites/site-a/offer/{DAY}/submitted", json={"by": "[TRADER]", "confirmation": "MDA-001"}, headers=OPERATOR)
    assert r.status_code == 200 and r.json()["status"] == "submitted" and r.json()["confirmation"] == "MDA-001"
    assert any(e["action"] == "offer.submitted" for e in store.audit)
    at(LATE)
    r = client.post(f"/v1/sites/site-a/offer/{DAY}/submitted", json={"by": "[TRADER]", "confirmation": "MDA-002"}, headers=OPERATOR)
    assert r.status_code == 409 and r.json()["error"]["code"] == "deadline_passed"
    r = client.post(f"/v1/sites/site-a/offer/{DAY}/submitted", json={"by": "[TRADER]", "confirmation": "MDA-002", "force": True}, headers=OPERATOR)
    assert r.status_code == 200 and r.json()["confirmation"] == "MDA-002"


def test_cleared_schedule_deviations_and_sample_day(client):
    approve_a(client)
    offered = [h["net_export_mw"] for h in client.get(f"/v1/sites/site-a/offer/{DAY}", headers=VIEWER).json()["hours"]]
    same = client.post(f"/v1/sites/site-a/cleared/{DAY}", json={"cleared_mw": offered}, headers=OPERATOR)
    assert same.status_code == 201 and same.json()["status"] == "as_offered" and same.json()["deviations"] == []
    cut = list(offered); cut[19] = offered[19] - 2
    r = client.post(f"/v1/sites/site-a/cleared/{DAY}", json={"cleared_mw": cut, "notes": [{"hour": 19, "reason": "congestion"}]}, headers=OPERATOR)
    j = r.json()
    assert j["status"] == "deviations" and len(j["deviations"]) == 1 and j["deviations"][0]["reason"] == "congestion" and j["margin_impact"] < 0
    assert client.get(f"/v1/sites/site-a/offer/{DAY}", headers=VIEWER).json()["status"] == "cleared"
    y = client.get("/v1/sites/site-a/cleared/2026-10-16", headers=VIEWER).json()
    assert [d["hour"] for d in y["deviations"]] == [22] and y["deviations"][0]["cleared_mw"] == 1.5
    rep = client.get("/v1/sites/site-a/replay/2026-10-16", headers=VIEWER).json()
    assert rep["hours"][22]["cleared_mw"] == 1.5 and rep["hours"][19]["cleared_mw"] == 10
    assert rep["sources"]["dispatched"] == "ems" and rep["sources"]["settlement"].startswith("preliminary statement 2026-10-17")


def test_imports_and_data_sources(client):
    meter = "hour,kwh_export,kwh_import\n" + "\n".join(f"{h},{1000 + h},{50}" for h in range(24)) + "\n"
    r = client.post("/v1/sites/site-a/imports/meter?date=2026-10-16", files={"file": ("m.csv", meter, "text/csv")}, headers=OPERATOR)
    assert r.status_code == 201 and r.json()["rows"] == 24
    bad = client.post("/v1/sites/site-a/imports/meter?date=2026-10-16", files={"file": ("m.csv", "hour,kwh\n1,2\n", "text/csv")}, headers=OPERATOR)
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "csv_columns"
    stmt = "date,energy_mxn,capacity_mxn,charges_mxn,kind\n2026-10-15,80100,0,1200,final\n2026-10-16,80400,0,1210,preliminary\n"
    r = client.post("/v1/sites/site-a/imports/statement", files={"file": ("s.csv", stmt, "text/csv")}, headers=OPERATOR)
    assert r.status_code == 201 and r.json()["rows"] == 2
    ds = client.get("/v1/sites/site-a/data-sources", headers=VIEWER).json()
    by = {s["name"]: s for s in ds["sources"]}
    assert by["meter"]["status"] == "live" and by["statements"]["status"] == "live" and "1 preliminary, 1 final" in by["statements"]["detail"]
    assert by["cenace_instructions"]["status"] == "pending" and "cenace_instructions" in ds["pending"]
    assert client.get("/v1/sites/site-a/replay/2026-10-16", headers=VIEWER).json()["sources"]["measured"] == "meter"
    g = client.get("/v1/sites/site-g/data-sources", headers=VIEWER).json()
    assert "registration" in g["blocked"]


def test_interconnection_export_cap_clips_offer(client, store):
    ic = client.get("/v1/sites/site-a/interconnection", headers=VIEWER).json()
    assert ic["export_cap_mw"] == 38.0 and ic["grid_charging"] == "allowed_with_limit"
    assert client.put("/v1/sites/site-a/interconnection", json={**ic, "export_cap_mw": 12}, headers=OPERATOR).status_code == 403
    r = client.put("/v1/sites/site-a/interconnection", json={**ic, "export_cap_mw": 12}, headers=ADMIN)
    assert r.status_code == 200 and any(e["action"] == "interconnection.update" for e in store.audit)
    plan = client.get(f"/v1/sites/site-a/plans/{DAY}", headers=VIEWER).json()
    assert all(h["solar_mw"] + h["dispatch_mw"] <= 12 + 1e-6 for h in plan["hours"] if h["dispatch_mw"] > 0)
    assert any(h["dispatch_mw"] > 0 for h in plan["hours"])
    approve_a(client)
    assert all(h["net_export_mw"] <= 12 + 1e-6 for h in client.get(f"/v1/sites/site-a/offer/{DAY}", headers=VIEWER).json()["hours"])


def test_battery_state_model_requires_signoff(client, store):
    a = client.get("/v1/sites/site-a/battery", headers=VIEWER).json()
    assert a["state_source"] == "ems" and a["approval_needs_signoff"] is False and a["soc_model"] is None
    c = client.get("/v1/sites/site-c/battery", headers=VIEWER).json()
    assert c["state_source"] == "model" and c["approval_needs_signoff"] is True and c["soc_model"]["band"][0] < c["soc_model"]["value"]
    r = client.post(f"/v1/sites/site-c/plans/{DAY}/approve", json={"strategy": "balanced"}, headers=OPERATOR)
    assert r.status_code == 428 and r.json()["error"]["code"] == "model_needs_signoff"
    r = client.post(f"/v1/sites/site-c/plans/{DAY}/approve", json={"strategy": "balanced", "sign_off": True}, headers=OPERATOR)
    assert r.status_code == 428
    r = client.post(f"/v1/sites/site-c/plans/{DAY}/approve", json={"strategy": "balanced", "sign_off": True, "acknowledge_soft_violations": True}, headers=MANAGER)
    assert r.status_code == 200, r.text
    assert "modelled battery state" in next(e["detail"] for e in store.audit if e["action"] == "plan.approve" and e["site_id"] == "site-c")


# ---------- product lines and business plan ----------

def test_products_list_and_detail(client):
    r = client.get("/v1/products", headers=VIEWER)
    assert r.status_code == 200
    lines = r.json()
    assert [l["number"] for l in lines] == list(range(1, 10))
    d = client.get("/v1/products/plant-forecasting", headers=VIEWER).json()
    assert d["sample"]["realized_mxn"] == 36700 and d["sample"]["theoretical_max_mxn"] == 47100
    es = client.get("/v1/products/critical-hours", headers={**VIEWER, **ES}).json()
    assert es["name"] == "Servicio de horas críticas" and "febrero" in es["sample"]["note"]
    assert client.get("/v1/products/nope", headers=VIEWER).status_code == 404


def test_business_plan_admin_only_and_adds_up(client):
    assert client.get("/v1/business/plan", headers=MANAGER).status_code == 403
    p = client.get("/v1/business/plan", headers=ADMIN).json()
    assert round(sum(l["arr_2030_usd"] for l in p["lines"])) == p["target_arr_usd"] == 33_800_000
    assert p["ramp"][-1]["arr_usd"] == 33_800_000 and p["lines"][0]["id"] == "battery-optimization"


def test_prospects_only_for_genau(client):
    assert client.get("/v1/products/white-label", headers=MANAGER).json()["prospects"] == []
    ps = client.get("/v1/products/white-label", headers=ADMIN).json()["prospects"]
    assert any(p["name"] == "SkySense" for p in ps) and all((p["url"] or "").startswith("https://") or p["source"] for p in ps)


# ---------- QEnergy offers and settlements ----------

import io as _io
import openpyxl as _openpyxl
from app.qenergy_store import QStore
from app.routers.qenergy import get_qstore

ENGINE, QENERGY = {"X-API-Key": "demo-engine"}, {"X-API-Key": "demo-qenergy"}


class FakeCalendar:
    """Stands in for an OpenTimestamps calendar: promises at once, 'confirms' in Bitcoin block 900000 once CONFIRMED."""
    CONFIRMED, DOWN = False, False

    def __init__(self, url):
        self.url = url

    def submit(self, digest, timeout=None):
        from opentimestamps.core.notary import PendingAttestation
        from opentimestamps.core.op import OpAppend, OpSHA256
        from opentimestamps.core.timestamp import Timestamp
        if FakeCalendar.DOWN:
            raise ConnectionError("calendar down")
        ts = Timestamp(digest)
        leaf = ts.ops.add(OpAppend(b"nonce-" + self.url.encode()[-8:]))
        commit = leaf.ops.add(OpSHA256())
        commit.attestations.add(PendingAttestation(self.url))
        return ts

    def get_timestamp(self, commitment, timeout=None):
        from opentimestamps.core.notary import BitcoinBlockHeaderAttestation
        from opentimestamps.core.timestamp import Timestamp
        if not FakeCalendar.CONFIRMED:
            raise KeyError("not yet in a block")
        ts = Timestamp(commitment)
        ts.attestations.add(BitcoinBlockHeaderAttestation(900000))
        return ts


class FakeRekor:
    """Stands in for rekor.sigstore.dev: checks Pantheon's signature over the hash, then logs it with its own clock."""
    entries, DOWN = [], False

    @staticmethod
    def post(body):
        import base64, time
        from cryptography.hazmat.primitives import hashes, serialization
        from cryptography.hazmat.primitives.asymmetric import ec, utils
        if FakeRekor.DOWN:
            raise ConnectionError("rekor down")
        spec = body["spec"]
        pub = serialization.load_pem_public_key(base64.b64decode(spec["signature"]["publicKey"]["content"]))
        digest = bytes.fromhex(spec["data"]["hash"]["value"])
        pub.verify(base64.b64decode(spec["signature"]["content"]), digest, ec.ECDSA(utils.Prehashed(hashes.SHA256())))
        FakeRekor.entries.append(spec["data"]["hash"]["value"])
        n = len(FakeRekor.entries)
        return {f"uuid{n}": {"logIndex": 1000 + n, "integratedTime": int(time.time())}}


@pytest.fixture
def qclient(client, tmp_path, monkeypatch):
    import opentimestamps.calendar
    monkeypatch.setattr(opentimestamps.calendar, "RemoteCalendar", FakeCalendar)
    FakeCalendar.CONFIRMED = FakeCalendar.DOWN = False
    from app.routers import vispera as vr
    FakeRekor.entries, FakeRekor.DOWN = [], False
    monkeypatch.setattr(vr, "_rekor_override", FakeRekor.post)
    monkeypatch.setenv("PANTHEON_VISPERA_ALLOW_EPHEMERAL", "true")
    qs = QStore(path=str(tmp_path / "q.db"), sample=True)
    app.dependency_overrides[get_qstore] = lambda: qs
    yield client


def _fc(day_mw=0.5):
    return {"model_version": "test-1", "hours": [
        {"hour": h, "sell_mw": day_mw if 10 <= h <= 16 else 0, "sell_price": 0, "buy_fixed_mw": 0.02,
         "buy_segments": [{"mw": 0.05, "price": 900}] if h in (3, 4) else []} for h in range(1, 25)]}


def test_engine_posts_and_partner_downloads_offers(qclient):
    c = qclient
    r = c.post("/v1/qenergy/sites/atotonilco/forecasts/2026-10-20", json=_fc(), headers=ENGINE)
    assert r.status_code == 201, r.text
    rec = r.json()
    assert rec["version"] == 1 and len(rec["sha256"]) == 64
    assert c.post("/v1/qenergy/sites/atotonilco/forecasts/2026-10-20", json=_fc(0.6), headers=ENGINE).json()["version"] == 2
    # partner cannot post forecasts; engine cannot read offers
    assert c.post("/v1/qenergy/sites/atotonilco/forecasts/2026-10-20", json=_fc(), headers=QENERGY).status_code == 403
    assert c.get("/v1/qenergy/sites/atotonilco/offers/2026-10-20/venta.xlsx", headers=ENGINE).status_code == 403
    # sales offer in QEnergy's template, latest version
    r = c.get("/v1/qenergy/sites/atotonilco/offers/2026-10-20/venta.xlsx", headers=QENERGY)
    assert r.status_code == 200 and r.headers["x-forecast-sha256"] != rec["sha256"]
    ws = _openpyxl.load_workbook(_io.BytesIO(r.content)).active
    assert ws["A2"].value == "ONDvr1" and ws["L4"].value == "01GTT" and ws["P4"].value == "01GTT-EX1"
    assert ws["D4"].value.date().isoformat() == "2026-10-20" and ws["D25"].value == 0.6 and ws["D14"].value == 0 and ws["D38"].value is None
    # purchase offer, 7 days from the date
    r = c.get("/v1/qenergy/sites/atotonilco/offers/2026-10-20/compra.xlsx", headers=QENERGY)
    ws = _openpyxl.load_workbook(_io.BytesIO(r.content)).active
    assert ws["A2"].value == "OCEvr1" and ws["M3"].value == "AXC0611GTT79" and ws["H3"].value == "=D3+6"
    assert ws["D16"].value == 0.02 and ws["F16"].value == 0.05 and ws["H16"].value == 900 and ws["J16"].value == 0
    r = c.get("/v1/qenergy/sites/atotonilco/offers/2026-10-20/compra.xlsx?days=1", headers=QENERGY)
    ws = _openpyxl.load_workbook(_io.BytesIO(r.content)).active
    assert ws["H3"].value.date().isoformat() == "2026-10-20"
    # history keeps both versions
    assert len(c.get("/v1/qenergy/sites/atotonilco/forecasts/2026-10-20/history", headers=QENERGY).json()) == 2
    # anyone can re-check the fingerprint from the canonical bytes
    import hashlib
    r = c.get("/v1/qenergy/sites/atotonilco/forecasts/2026-10-20/canonical", headers=QENERGY)
    assert hashlib.sha256(r.content).hexdigest() == r.headers["x-forecast-sha256"]
    assert b'"model_version":"test-1"' in r.content and b'"day":"2026-10-20"' in r.content


def test_forecast_validation_and_missing(qclient):
    bad = _fc(); bad["hours"] = bad["hours"][:23]
    assert qclient.post("/v1/qenergy/sites/atotonilco/forecasts/2026-10-21", json=bad, headers=ENGINE).status_code == 422
    assert qclient.get("/v1/qenergy/sites/atotonilco/offers/2030-01-01/venta.xlsx", headers=QENERGY).status_code == 404
    assert qclient.get("/v1/qenergy/sites/nowhere/forecasts", headers=QENERGY).status_code == 404


def test_settlement_math(qclient):
    c = qclient
    c.post("/v1/qenergy/sites/atotonilco/forecasts/2026-11-02", json=_fc(), headers=ENGINE)
    hours = [{"hour": h, "cleared_mw": 1.0 if h == 12 else 0, "measured_mwh": 0.8 if h == 12 else 0, "pml_mda": 2000, "pml_mtr": 2500} for h in range(1, 25)]
    assert c.post("/v1/qenergy/sites/atotonilco/results/2026-11-02", json={"hours": hours}, headers=QENERGY).status_code == 201
    s = c.get("/v1/qenergy/sites/atotonilco/settlements?month=2026-11", headers=QENERGY).json()
    d = s["days"][0]
    assert d["mda_mxn"] == 2000 and d["mtr_mxn"] == -500 and d["total_mxn"] == 1500 and not s["sample_data"]
    assert c.get("/v1/qenergy/sites/atotonilco/settlements?month=2026-11&format=xlsx", headers=QENERGY).status_code == 200
    assert "mda_mxn" in c.get("/v1/qenergy/sites/atotonilco/settlements?month=2026-11&format=csv", headers=QENERGY).text
    # sample month exists for the page demo
    assert len(c.get("/v1/qenergy/sites/atotonilco/settlements?month=2026-10", headers=QENERGY).json()["days"]) >= 1



# ---------- Víspera chain ----------
ADMIN, CLIENT = {"X-API-Key": "demo-admin"}, {"X-API-Key": "demo-client"}
TERMS = {"chain_id": "atotonilco", "subject": {"site": "atotonilco", "central": "01GTT"}, "quantity": "hourly sales offer, MW",
         "start_date": "2030-01-01", "scoring": "MAE of offered vs metered MWh, every day from the start date"}


def _start(c, **kw):
    return c.post("/v1/vispera/chains", json={**TERMS, **kw}, headers=ADMIN)


def test_chain_terms_forecasts_actuals_and_public_view(qclient):
    c = qclient
    assert c.post("/v1/vispera/chains", json=TERMS, headers=ENGINE).status_code == 403
    assert _start(c, start_date="2020-01-01").status_code == 422          # terms must come before the start date
    r = _start(c)
    assert r.status_code == 201 and r.json()["seq"] == 0 and r.json()["kind"] == "terms"
    assert _start(c).status_code == 409
    # before the start date: recorded for QEnergy but not chained
    assert c.post("/v1/qenergy/sites/atotonilco/forecasts/2029-12-31", json=_fc(), headers=ENGINE).json()["chain"] is None
    r1 = c.post("/v1/qenergy/sites/atotonilco/forecasts/2030-01-01", json=_fc(0.4), headers=ENGINE).json()
    r2 = c.post("/v1/qenergy/sites/atotonilco/forecasts/2030-01-01", json=_fc(0.5), headers=ENGINE).json()
    r3 = c.post("/v1/qenergy/sites/atotonilco/forecasts/2030-01-02", json=_fc(0.6), headers=ENGINE).json()
    assert [r1["chain"]["seq"], r2["chain"]["seq"], r3["chain"]["seq"]] == [1, 2, 3]
    res = {"hours": [{"hour": h, "cleared_mw": 0.4, "measured_mwh": 0.39, "pml_mda": 1500 + h, "pml_mtr": 1600} for h in range(1, 25)],
           "vintage": "MDA 2030-01-01 v1", "retrieved_at": "2030-01-01T18:30:00-06:00"}
    assert c.post("/v1/qenergy/sites/atotonilco/results/2030-01-01", json=res, headers=QENERGY).json()["chain"]["seq"] == 4
    # anyone, no key: hashes and anchors, never a forecast
    pub = c.get("/v1/vispera/chains/atotonilco").json()
    assert pub["length"] == 5 and pub["summary"]["forecasts"] == 3 and pub["summary"]["forecast_days"] == 2
    assert all(e["content"] is None for e in pub["entries"] if e["kind"] == "forecast")
    assert "0.4" not in json.dumps(pub["entries"][1])
    act = json.loads(pub["entries"][4]["content"])
    assert act["vintage"] == "MDA 2030-01-01 v1" and "measured_mwh" not in pub["entries"][4]["content"]  # QEnergy volumes stay private
    assert pub["summary"]["rekor"] == 5 and pub["summary"]["rekor_before_cutoff"] == 3
    assert all(e["bitcoin"]["status"] == "pending" for e in pub["entries"])
    assert c.get(pub["entries"][1]["bitcoin"]["proof_url"]).status_code == 200
    assert c.get("/v1/vispera/index").json()["chains"][0]["length"] == 5
    assert "BEGIN PUBLIC KEY" in c.get("/v1/vispera/public-key").text
    # the forecast page links to its chain entry
    assert c.get("/v1/qenergy/sites/atotonilco/forecasts/2030-01-02", headers=QENERGY).json()["chain"]["seq"] == 3


def test_client_disclosure_verifies_and_tampering_fails(qclient, tmp_path):
    import copy, importlib.util
    from app import vispera as V
    c = qclient
    _start(c)
    for d, mw in [("2030-01-01", 0.4), ("2030-01-02", 0.5), ("2030-02-01", 0.6)]:  # two months, two chain keys
        c.post(f"/v1/qenergy/sites/atotonilco/forecasts/{d}", json=_fc(mw), headers=ENGINE)
    assert c.get("/v1/vispera/chains/atotonilco/disclosure", headers=ENGINE).status_code == 403
    d = c.get("/v1/vispera/chains/atotonilco/disclosure", headers=CLIENT).json()
    assert set(d["chain_keys"]) == {"2030-01", "2030-02"} and d["entries"][0]["seq"] == 0
    assert V.verify_chain(d["entries"], d["chain_keys"])["ok"]
    spec = importlib.util.spec_from_file_location("vv", "tools/verify_vispera.py"); vv = importlib.util.module_from_spec(spec); spec.loader.exec_module(vv)
    assert vv.check(d, "client", online=False) == 0
    # change one number in a forecast
    t = copy.deepcopy(d); t["entries"][2]["content"] = t["entries"][2]["content"].replace('"sell_mw":0.5', '"sell_mw":0.501')
    assert vv.check(t, "client", online=False) == 1
    # drop a bad day
    t = copy.deepcopy(d); del t["entries"][2]
    assert vv.check(t, "client", online=False) == 1 and vv.check(t, "public", online=False) == 1
    # relabel a forecast to another day
    t = copy.deepcopy(d); t["entries"][1]["target_day"] = "2030-01-05"
    assert vv.check(t, "client", online=False) == 1
    # without the right month key the salts can't be derived
    t = copy.deepcopy(d); t["chain_keys"]["2030-02"] = "00" * 32
    assert not V.verify_chain(t["entries"], t["chain_keys"])["ok"]
    # the public chain alone checks numbering, links and cutoffs
    assert vv.check(c.get("/v1/vispera/chains/atotonilco").json(), "public", online=False) == 0
    # the escrowed month key yields the same chain key the client received
    keys = {k["month"]: k["key"] for k in c.get("/v1/vispera/keys", headers=ADMIN).json()["keys"]}
    assert V.chain_key(bytes.fromhex(keys["2030-01"]), "atotonilco").hex() == d["chain_keys"]["2030-01"]
    assert c.get("/v1/vispera/keys", headers=CLIENT).status_code == 403


def test_anchor_retries_and_bitcoin_upgrade(qclient):
    c = qclient
    FakeRekor.DOWN = FakeCalendar.DOWN = True
    _start(c)
    c.post("/v1/qenergy/sites/atotonilco/forecasts/2030-01-01", json=_fc(), headers=ENGINE)
    pub = c.get("/v1/vispera/chains/atotonilco").json()
    assert pub["summary"]["rekor"] == 0 and pub["entries"][1]["bitcoin"]["status"] == "failed"
    FakeRekor.DOWN = FakeCalendar.DOWN = False
    out = c.post("/v1/vispera/refresh", headers=ENGINE).json()
    assert out["stamped"] == 2 and out["rekor_logged"] == 2
    FakeCalendar.CONFIRMED = True
    assert c.post("/v1/vispera/refresh", headers=ENGINE).json()["anchored"] == 2
    pub = c.get("/v1/vispera/chains/atotonilco").json()
    assert all(e["bitcoin"]["block_height"] == 900000 for e in pub["entries"]) and pub["summary"]["rekor"] == 2


def test_chain_needs_a_lasting_key(qclient, monkeypatch):
    monkeypatch.delenv("PANTHEON_VISPERA_ALLOW_EPHEMERAL")
    assert _start(qclient).status_code == 409
